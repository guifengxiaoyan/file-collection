# -*- coding: utf-8 -*-
"""针对 2026-10-16 测试报告所发现问题的修复验证测试（test_fixes.py）。

覆盖六个修复点：
  X1. 创建/编辑主题缺少或非法 deadline（title）时返回表单提示而非 500
  X2. Excel 导入跳过表头行；创建主题时文本/Excel/模板三来源跨来源去重
  X3. “合并相同的行”改为跨所有文件的全局去重（不限相邻行）
  X4. 中文下载名的 ASCII 回退名可读（不再丢成 ".txt"），filename* 保留
  X5. 已完成/已过期的上传页 JS 无空指针报错（真实 Chrome 验证）
  X6. 仓库卫生：.gitignore/requirements.txt/Dockerfile（静态检查）
"""
import io
import os
import re
import sys
import time
import zipfile
import sqlite3
from datetime import datetime, timedelta

import requests
from openpyxl import Workbook, load_workbook

BASE = 'http://127.0.0.1:5000'
HERE = os.path.dirname(os.path.abspath(__file__))
FIX = os.path.join(HERE, 'fixtures')
DB = os.path.join(HERE, 'instance', 'file_collection.db')
CHROME = r'C:\Program Files\Google\Chrome\Application\chrome.exe'

PASS, FAIL = [], []


def check(name, cond, detail=''):
    if cond:
        PASS.append(name)
        print(f'  ✅ {name}')
    else:
        FAIL.append((name, detail))
        print(f'  ❌ {name}  {detail}')


def db_query(sql, args=()):
    con = sqlite3.connect(DB)
    try:
        return con.execute(sql, args).fetchall()
    finally:
        con.close()


def reset_data():
    """清空业务数据与上传文件（保留 admin），保证测试可重复。"""
    import shutil
    con = sqlite3.connect(DB)
    try:
        for t in ['announcement_attachment', 'announcement', 'attachment',
                  'collection_object', 'collection_object_template',
                  'theme_attachment', 'collection_theme']:
            con.execute(f'DELETE FROM {t}')
        con.commit()
    finally:
        con.close()
    up = os.path.join(HERE, 'uploads')
    if os.path.isdir(up):
        for name in os.listdir(up):
            p = os.path.join(up, name)
            try:
                if os.path.isdir(p):
                    shutil.rmtree(p)
                else:
                    os.remove(p)
            except OSError:
                pass


def wait(sec=0.3):
    time.sleep(sec)


def fx(fname):
    with open(os.path.join(FIX, fname), 'rb') as f:
        return f.read()


def upload_file(fname, content=None):
    return (fname, content if content is not None else fx(fname))


reset_data()

# ---------- 生成本套件需要的 Excel fixture ----------
def _make_names_excel(path, first_cell, names):
    wb = Workbook()
    ws = wb.active
    ws.title = '名单'
    if first_cell is not None:
        ws.cell(row=1, column=1, value=first_cell)
    start = 2 if first_cell is not None else 1
    for i, n in enumerate(names, start):
        ws.cell(row=i, column=1, value=n)
    wb.save(path)


_make_names_excel(os.path.join(FIX, 'objects_noheader.xlsx'), None, ['甲一', '甲二', '甲三'])
_make_names_excel(os.path.join(FIX, 'objects_composite_header.xlsx'), '部门名称',
                  ['教务处', '财务处', '后勤处'])
print('fixtures: objects_noheader.xlsx / objects_composite_header.xlsx 已生成')

s = requests.Session()
s.trust_env = False  # 绕过沙箱代理，直连本地服务

r = s.post(BASE + '/admin/login', data={'username': 'admin', 'password': 'admin123'},
           allow_redirects=False)
assert r.status_code == 302, f'登录失败: {r.status_code}'
print('已登录 admin')

future = (datetime.now() + timedelta(days=7)).strftime('%Y-%m-%dT%H:%M')
past = (datetime.now() - timedelta(days=1)).strftime('%Y-%m-%dT%H:%M')

# =====================================================================
print()
print('=' * 70)
print('X1. 截止时间/标题的服务端校验（问题2：缺字段不再 500）')
print('=' * 70)

before = db_query('SELECT COUNT(*) FROM collection_theme')[0][0]

r = s.post(BASE + '/admin/theme/create',
           data={'title': '无截止时间主题', 'collector_objects': '甲'},
           allow_redirects=False)
check('X1.1 缺 deadline 返回 200 表单页（而非 500）', r.status_code == 200, f'code={r.status_code}')
check('X1.2 缺 deadline 有错误提示', '截止时间' in r.text)
check('X1.3 缺 deadline 不创建主题', db_query('SELECT COUNT(*) FROM collection_theme')[0][0] == before)

r = s.post(BASE + '/admin/theme/create',
           data={'title': '非法截止时间主题', 'deadline': 'not-a-date'},
           allow_redirects=False)
check('X1.4 非法 deadline 返回 200', r.status_code == 200, f'code={r.status_code}')
check('X1.5 非法 deadline 不创建主题', db_query('SELECT COUNT(*) FROM collection_theme')[0][0] == before)

r = s.post(BASE + '/admin/theme/create',
           data={'title': '  ', 'deadline': future},
           allow_redirects=False)
check('X1.6 空标题返回 200 且有提示', r.status_code == 200 and '主题名称' in r.text)
check('X1.7 空标题不创建主题', db_query('SELECT COUNT(*) FROM collection_theme')[0][0] == before)

# 正常创建仍然成功（回归）
r = s.post(BASE + '/admin/theme/create',
           data={'title': 'X1回归主题', 'deadline': future, 'collector_objects': '回归甲'},
           allow_redirects=False)
check('X1.8 合法参数仍可正常创建（302）', r.status_code == 302, f'code={r.status_code}')
x1_theme_id = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
old_deadline = db_query('SELECT deadline FROM collection_theme WHERE id=?', (x1_theme_id,))[0][0]

# edit_theme 校验
r = s.post(BASE + f'/admin/theme/{x1_theme_id}/edit',
           data={'title': 'X1回归主题', 'deadline': ''},
           allow_redirects=False)
check('X1.9 编辑缺 deadline 返回 200', r.status_code == 200, f'code={r.status_code}')
check('X1.10 编辑缺 deadline 不修改原截止时间',
      db_query('SELECT deadline FROM collection_theme WHERE id=?', (x1_theme_id,))[0][0] == old_deadline)

r = s.post(BASE + f'/admin/theme/{x1_theme_id}/edit',
           data={'title': 'X1回归主题', 'deadline': '2026/13/99'},
           allow_redirects=False)
check('X1.11 编辑非法 deadline 返回 200', r.status_code == 200, f'code={r.status_code}')
check('X1.12 编辑非法 deadline 不修改原截止时间',
      db_query('SELECT deadline FROM collection_theme WHERE id=?', (x1_theme_id,))[0][0] == old_deadline)

r = s.post(BASE + f'/admin/theme/{x1_theme_id}/edit',
           data={'title': 'X1回归主题改名', 'deadline': future},
           allow_redirects=False)
check('X1.13 合法编辑仍成功（302）', r.status_code == 302, f'code={r.status_code}')
check('X1.14 编辑生效', db_query('SELECT title FROM collection_theme WHERE id=?', (x1_theme_id,))[0][0] == 'X1回归主题改名')

# =====================================================================
print()
print('=' * 70)
print('X2. Excel 导入跳表头 + 三来源跨来源去重（问题3）')
print('=' * 70)


def create_theme(title, objects_text='', excel=None, templates=None):
    data = {'title': title, 'deadline': future}
    if objects_text:
        data['collector_objects'] = objects_text
    files = {}
    if excel:
        files['objects_excel_file'] = excel
    if templates:
        # requests 多值字段
        data['selected_templates'] = templates
    r = s.post(BASE + '/admin/theme/create', data=data, files=files, allow_redirects=False)
    return r


def theme_objects(theme_id):
    return [x[0] for x in db_query(
        'SELECT name FROM collection_object WHERE theme_id=? ORDER BY id', (theme_id,))]


# 2.1 带表头 Excel：表头不导入，数据全导入
r = create_theme('X2表头主题', excel=upload_file('objects.xlsx'))
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
objs = theme_objects(tid)
check('X2.1 带表头 Excel：导入 5 个对象', len(objs) == 5, f'objs={objs}')
check('X2.2 表头“姓名”不再是收集对象', '姓名' not in objs, f'objs={objs}')

# 2.3 无表头 Excel：首行即数据，全部导入
r = create_theme('X2无表头主题', excel=upload_file('objects_noheader.xlsx'))
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
objs = theme_objects(tid)
check('X2.3 无表头 Excel：首行数据不被跳过（3 个）', objs == ['甲一', '甲二', '甲三'], f'objs={objs}')

# 2.4 复合表头（部门名称）被识别并跳过
r = create_theme('X2复合表头主题', excel=upload_file('objects_composite_header.xlsx'))
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
objs = theme_objects(tid)
check('X2.4 复合表头被跳过（3 个对象）', objs == ['教务处', '财务处', '后勤处'], f'objs={objs}')

# 2.5 文本 + Excel 重叠 -> 跨来源去重
r = create_theme('X2去重主题', objects_text='张三\n李四', excel=upload_file('objects.xlsx'))
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
objs = theme_objects(tid)
check('X2.5 文本+Excel 跨来源去重（5 个，无重复）',
      objs == ['张三', '李四', '王五', '赵六', '钱七'], f'objs={objs}')

# 2.6 文本框内重复行也去重
r = create_theme('X2文本去重主题', objects_text='张三\n张三\n李四\n\n  \n李四')
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
objs = theme_objects(tid)
check('X2.6 文本框内重复/空行去重（2 个）', objs == ['张三', '李四'], f'objs={objs}')

# 2.7 模板 + 文本重叠 -> 去重
s.post(BASE + '/admin/object-templates', data={'names': '张三', 'category': 'X2测试分组'})
wait()
tpl_row = db_query("SELECT id FROM collection_object_template WHERE name='张三' ORDER BY id DESC LIMIT 1")
if tpl_row:
    tpl_id = tpl_row[0][0]
    r = create_theme('X2模板去重主题', objects_text='张三\n李四', templates=[str(tpl_id)])
    tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
    objs = theme_objects(tid)
    check('X2.7 模板与文本去重（2 个）', objs == ['张三', '李四'], f'objs={objs}')
else:
    check('X2.7 模板与文本去重（2 个）', False, '模板创建失败')

# 2.8 管理端 Excel 导入：表头不导入 + 按名称去重
r = create_theme('X2管理端主题', objects_text='张三\n李四')
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
r = s.post(BASE + f'/admin/theme/{tid}/objects',
           data={'import_excel': '1'}, files={'excel_file': upload_file('objects.xlsx')},
           allow_redirects=False)
wait()
objs = theme_objects(tid)
check('X2.8 管理端导入跳表头（5 个，无“姓名”）',
      len(objs) == 5 and '姓名' not in objs, f'objs={objs}')
# 该路由 POST 后直接渲染页面（非 302），flash 在本次响应中就地消费
check('X2.9 导入结果提示“共导入 3 条数据”', '共导入 3 条数据' in r.text)

# 2.10 损坏的 Excel 不再 500
r = s.post(BASE + f'/admin/theme/{tid}/objects',
           data={'import_excel': '1'},
           files={'excel_file': ('坏文件.xlsx', b'this is not an excel file')},
           allow_redirects=False)
check('X2.10 损坏 Excel 返回 200 且有提示（而非 500）',
      r.status_code == 200 and '解析失败' in r.text, f'code={r.status_code}')

# =====================================================================
print()
print('=' * 70)
print('X3. 合并去重改为全局去重（问题4：不限相邻行）')
print('=' * 70)

sys.path.insert(0, HERE)
from utils import merge_excel_files  # noqa: E402

files_args = [
    {'path': os.path.join(FIX, 'submit_张三.xlsx'), 'object_name': '张三', 'original_name': 'submit_张三.xlsx'},
    {'path': os.path.join(FIX, 'submit_李四.xlsx'), 'object_name': '李四', 'original_name': 'submit_李四.xlsx'},
]

# 3.1 不开去重：5 行（回归）
buf = merge_excel_files(files_args, merge_identical=False)
wb = load_workbook(io.BytesIO(buf.getvalue()))
ws = wb.active
names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
check('X3.1 不去重时 5 行（回归）', names == ['设备A', '设备B', '设备C', '设备A', '设备D'], f'names={names}')

# 3.2 全局去重：设备A 在张三文件第 2 行、李四文件第 2 行（不相邻）应被合并
buf = merge_excel_files(files_args, merge_identical=True)
wb = load_workbook(io.BytesIO(buf.getvalue()))
ws = wb.active
names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
check('X3.2 全局去重：不相邻重复行被合并（5→4）',
      names == ['设备A', '设备B', '设备C', '设备D'], f'names={names}')

# 3.3 被去重行（李四的设备A 行）的图片固定到保留行（输出第 2 行，0-based row 1）
img_rows = sorted(im.anchor._from.row for im in ws._images)
img_cols = sorted(im.anchor._from.col for im in ws._images)
check('X3.3 去重后图片不错位（2 张图都在保留行 row=1）', img_rows == [1, 1], f'rows={img_rows}')
# 张三图 anchor_col=3、李四图 anchor_col=2（均为 0-based 列），去重后列保持不变
check('X3.4 图片列保持原锚定（col 2 与 3）', img_cols == [2, 3], f'cols={img_cols}')

# 3.5 去重 + 来源列
buf = merge_excel_files(files_args, merge_identical=True, mark_source=True)
wb = load_workbook(io.BytesIO(buf.getvalue()))
ws = wb.active
header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
check('X3.5 去重+来源列：表头含“收集对象”', header[-1] == '收集对象', f'header={header}')
srcs = [ws.cell(row=r, column=ws.max_column).value for r in range(2, ws.max_row + 1)]
check('X3.6 去重+来源列：来源标注正确', srcs == ['张三', '张三', '张三', '李四'], f'srcs={srcs}')

# 3.7 HTTP 级验证：完整上传后走合并路由
r = create_theme('X3合并主题', objects_text='张三\n李四')
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
oid_a = db_query("SELECT id FROM collection_object WHERE theme_id=? AND name='张三'", (tid,))[0][0]
oid_b = db_query("SELECT id FROM collection_object WHERE theme_id=? AND name='李四'", (tid,))[0][0]
s.post(BASE + f'/upload/{oid_a}', files=[('file', upload_file('submit_张三.xlsx'))], allow_redirects=False)
s.post(BASE + f'/upload/{oid_b}', files=[('file', upload_file('submit_李四.xlsx'))], allow_redirects=False)
wait()
r = s.post(BASE + f'/admin/theme/{tid}/merge-excel',
           data={'merge_identical': 'on', 'mark_source': 'on'})
check('X3.7 合并路由返回 xlsx', r.status_code == 200 and r.content[:2] == b'PK', f'code={r.status_code}')
if r.content[:2] == b'PK':
    wb = load_workbook(io.BytesIO(r.content))
    ws = wb.active
    names = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    check('X3.8 HTTP 合并全局去重生效（4 行）',
          names == ['设备A', '设备B', '设备C', '设备D'], f'names={names}')
    check('X3.9 合并表格下载名含 filename*',
          "filename*=UTF-8''" in r.headers.get('Content-Disposition', ''),
          r.headers.get('Content-Disposition', ''))

# =====================================================================
print()
print('=' * 70)
print('X4. 中文下载名的 ASCII 回退（问题5）')
print('=' * 70)


def cd_fallback(resp):
    """从 Content-Disposition 中取出 ASCII 回退名 filename（兼容带/不带引号）。"""
    cd = resp.headers.get('Content-Disposition', '')
    m = re.search(r'filename="([^"]*)"', cd)
    if m:
        return m.group(1)
    m = re.search(r'filename=([^;]+)', cd)  # 纯 ASCII 名时 Werkzeug 不加引号
    return m.group(1).strip() if m else None


# 4.1 主题附件：中文原名下载
r = create_theme('X4下载主题', objects_text='测试员')
tid = db_query('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1')[0][0]
s.post(BASE + f'/admin/theme/{tid}/edit',
       files=[('attachments', ('张三.txt', 'zhangsan content'.encode()))],
       data={'title': 'X4下载主题', 'deadline': future}, allow_redirects=False)
wait()
att_id = db_query('SELECT id FROM theme_attachment WHERE theme_id=? ORDER BY id DESC LIMIT 1', (tid,))[0][0]
r = s.get(BASE + f'/theme/download/{att_id}')
cd = r.headers.get('Content-Disposition', '')
check('X4.1 中文附件下载 200', r.status_code == 200, f'code={r.status_code}')
check('X4.2 filename* 正确编码中文名', "filename*=UTF-8''%E5%BC%A0%E4%B8%89.txt" in cd, cd)
check('X4.3 ASCII 回退名可读（file.txt 而非 .txt）', cd_fallback(r) == 'file.txt', cd)

# 4.4 混合中英文文件名：保留 ASCII 部分
s.post(BASE + f'/admin/theme/{tid}/edit',
       files=[('attachments', ('汇总表 2026（最终版）.xlsx', b'fake xlsx'))],
       data={'title': 'X4下载主题', 'deadline': future}, allow_redirects=False)
wait()
att_id = db_query('SELECT id FROM theme_attachment WHERE theme_id=? ORDER BY id DESC LIMIT 1', (tid,))[0][0]
r = s.get(BASE + f'/theme/download/{att_id}')
check('X4.4 中英混合名回退保留 ASCII（2026.xlsx）', cd_fallback(r) == '2026.xlsx',
      r.headers.get('Content-Disposition', ''))

# 4.5 纯 ASCII 名不受影响
s.post(BASE + f'/admin/theme/{tid}/edit',
       files=[('attachments', ('readme.txt', b'ascii name'))],
       data={'title': 'X4下载主题', 'deadline': future}, allow_redirects=False)
wait()
att_id = db_query('SELECT id FROM theme_attachment WHERE theme_id=? ORDER BY id DESC LIMIT 1', (tid,))[0][0]
r = s.get(BASE + f'/theme/download/{att_id}')
check('X4.5 纯 ASCII 名下载名不变（readme.txt）', cd_fallback(r) == 'readme.txt',
      r.headers.get('Content-Disposition', ''))

# 4.6 公告附件同样处理
s.post(BASE + '/admin/announcement/create',
       data={'title': 'X4公告', 'content': 'c'},
       files=[('attachments', ('公告附件.pdf', b'%PDF fake'))], allow_redirects=False)
wait()
ann_att = db_query('SELECT id FROM announcement_attachment ORDER BY id DESC LIMIT 1')[0][0]
r = s.get(BASE + f'/announcement/download/{ann_att}')
check('X4.6 公告中文附件回退名可读（file.pdf）', cd_fallback(r) == 'file.pdf',
      r.headers.get('Content-Disposition', ''))

# 4.7 收集附件（按对象名命名）同样处理
oid = db_query("SELECT id FROM collection_object WHERE theme_id=? ORDER BY id LIMIT 1", (tid,))[0][0]
s.post(BASE + f'/upload/{oid}', files=[('file', ('李四的报告.pdf', b'%PDF fake'))], allow_redirects=False)
wait()
att_id = db_query('SELECT id FROM attachment ORDER BY id DESC LIMIT 1')[0][0]
r = s.get(BASE + f'/download/attachment/{att_id}')
check('X4.7 收集附件按对象名下载，回退名可读（file.pdf）', cd_fallback(r) == 'file.pdf',
      r.headers.get('Content-Disposition', ''))
check('X4.8 filename* 使用对象名（李四的报告.pdf）',
      "filename*=UTF-8''%E6%9D%8E%E5%9B%9B%E7%9A%84%E6%8A%A5%E5%91%8A.pdf" in
      r.headers.get('Content-Disposition', ''), r.headers.get('Content-Disposition', ''))

# 4.9 导出 zip（中文主题名）
r = s.get(BASE + f'/admin/theme/{tid}/export')
check('X4.9 导出 zip 200', r.status_code == 200 and r.content[:2] == b'PK', f'code={r.status_code}')
cd = r.headers.get('Content-Disposition', '')
check('X4.10 导出 zip 下载名含 filename* 且回退非空',
      "filename*=UTF-8''" in cd and cd_fallback(r) not in (None, '', '.zip'), cd)

# =====================================================================
print()
print('=' * 70)
print('X5. 已完成/已过期上传页 JS 无空指针（问题1，真实 Chrome）')
print('=' * 70)

# 已完成对象
s.post(BASE + f'/upload/{oid}', data={'finish': '1'}, allow_redirects=False)
wait()
# 已过期主题 + 对象
s.post(BASE + '/admin/theme/create',
       data={'title': 'X5过期主题', 'deadline': past, 'collector_objects': '过期人员'})
wait()
expired_obj = db_query("SELECT id FROM collection_object WHERE name='过期人员' ORDER BY id DESC LIMIT 1")[0][0]

js_errors = []
try:
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = p.chromium.launch(executable_path=CHROME, headless=True,
                                    args=['--no-sandbox', '--no-proxy-server'])
        page = browser.new_page()
        page.on('pageerror', lambda e: js_errors.append(str(e)))
        page.goto(BASE + f'/upload/{oid}', wait_until='networkidle')
        page.wait_for_timeout(400)
        check('X5.1 已完成对象上传页渲染完成状态', '您已完成文件上传' in page.content())
        page.goto(BASE + f'/upload/{expired_obj}', wait_until='networkidle')
        page.wait_for_timeout(400)
        check('X5.2 已过期上传页渲染截止提示', '该主题已超过截止时间' in page.content())
        # 完成按钮在两种状态下都应保持禁用
        page.goto(BASE + f'/upload/{oid}', wait_until='networkidle')
        page.wait_for_timeout(300)
        disabled = page.eval_on_selector('#finish-btn', 'el => el.disabled')
        check('X5.3 已完成状态“完成上传”按钮保持禁用', disabled is True)
        browser.close()
    check('X5.4 两个状态页均无 JS 未捕获异常', not js_errors, f'errors={js_errors}')
except ImportError:
    check('X5.4 两个状态页均无 JS 未捕获异常', False, 'playwright 不可用')

# =====================================================================
print()
print('=' * 70)
print('X6. 仓库卫生（静态检查）')
print('=' * 70)

with open(os.path.join(HERE, '.gitignore'), encoding='utf-8') as f:
    gi = f.read()
check('X6.1 .gitignore 已忽略 instance/', 'instance/' in gi)

with open(os.path.join(HERE, 'requirements.txt'), encoding='utf-8') as f:
    req = f.read()
check('X6.2 Pillow 版本已放宽（支持 3.13 wheel）', 'Pillow>=10.4' in req and 'Pillow==10.2.0' not in req)

with open(os.path.join(HERE, 'Dockerfile'), encoding='utf-8') as f:
    df = f.read()
check('X6.3 Dockerfile 无空 apt-get 安装', 'apt-get install -y --no-install-recommends \\\n    &&' not in df)

print()
print('=' * 70)
print(f'修复验证结果：通过 {len(PASS)} 项，失败 {len(FAIL)} 项')
if FAIL:
    for name, detail in FAIL:
        print(f'  - {name}: {detail}')
print('=' * 70)
sys.exit(1 if FAIL else 0)
