# -*- coding: utf-8 -*-
"""Excel 合并功能深度测试：条件检测、按行合并、去重、来源列、图片锚定、WPS DISPIMG 截图保留。"""
import io
import os
import re
import sys
import json
import zipfile
from datetime import datetime, timedelta

import requests
from openpyxl import Workbook, load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker, TwoCellAnchor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from utils import _inject_wps_cell_images  # 仅用于构造测试用 WPS 源文件

BASE = 'http://127.0.0.1:5000'
FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')
PASS, FAIL = [], []

def check(name, cond, detail=''):
    if cond:
        PASS.append(name)
        print(f'  ✅ {name}')
    else:
        FAIL.append((name, detail))
        print(f'  ❌ {name}  {detail}')

def fx(fname):
    with open(os.path.join(FIX, fname), 'rb') as f:
        return f.read()

s = requests.Session()
s.trust_env = False  # 绕过沙箱代理，直连本地服务
s.post(BASE + '/admin/login', data={'username': 'admin', 'password': 'admin123'})

def create_theme(title, objects):
    deadline = (datetime.now() + timedelta(days=7)).strftime('%Y-%m-%dT%H:%M')
    r = s.post(BASE + '/admin/theme/create',
               data={'title': title, 'deadline': deadline, 'collector_objects': '\n'.join(objects)},
               allow_redirects=False)
    assert r.status_code == 302, r.text
    import sqlite3
    con = sqlite3.connect(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'instance', 'file_collection.db'))
    tid = con.execute('SELECT id FROM collection_theme ORDER BY id DESC LIMIT 1').fetchone()[0]
    oids = {name: con.execute('SELECT id FROM collection_object WHERE theme_id=? AND name=?', (tid, name)).fetchone()[0]
            for name in objects}
    con.close()
    return tid, oids

def upload(object_id, fname, content=None):
    files = {'file': (fname, content if content is not None else fx(fname))}
    r = s.post(BASE + f'/upload/{object_id}', files=files, allow_redirects=False)
    assert r.status_code == 302, r.text

print('=' * 70)
print('M1. 合并条件检测')
print('=' * 70)

tid1, oids1 = create_theme('合并检测主题A', ['张三', '李四', '王五'])
upload(oids1['张三'], 'submit_张三.xlsx')
upload(oids1['李四'], 'submit_李四.xlsx')
upload(oids1['王五'], '说明.txt')
r = s.get(BASE + f'/admin/theme/{tid1}/merge-excel')
check('M1.1 混合非 Excel → 不可合并并给出原因', '暂不能合并' in r.text and '说明.txt' in r.text, r.text[:200])

# 删掉 txt，换成 6 列文件
import sqlite3
con = sqlite3.connect(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'instance', 'file_collection.db'))
att = con.execute('SELECT id FROM attachment WHERE collection_object_id=?', (oids1['王五'],)).fetchone()[0]
con.close()
s.post(BASE + f'/upload/delete-attachment/{att}')
upload(oids1['王五'], 'submit_王五_6cols.xlsx')
r = s.get(BASE + f'/admin/theme/{tid1}/merge-excel')
check('M1.2 列数不一致 → 不可合并并说明', '暂不能合并' in r.text and '列数' in r.text, r.text[:300])

# 空主题（无任何 Excel）
tid2, oids2 = create_theme('合并检测主题B', ['测试员'])
upload(oids2['测试员'], '说明.txt')
r = s.get(BASE + f'/admin/theme/{tid2}/merge-excel')
check('M1.3 无 Excel 附件 → 不可合并', '暂不能合并' in r.text and '没有任何 Excel 附件' in r.text)

print()
print('=' * 70)
print('M2. 满足条件的合并（按行堆叠 + 图片保留）')
print('=' * 70)

tid3, oids3 = create_theme('合并执行主题', ['张三', '李四'])
upload(oids3['张三'], 'submit_张三.xlsx')
upload(oids3['李四'], 'submit_李四.xlsx')
r = s.get(BASE + f'/admin/theme/{tid3}/merge-excel')
check('M2.1 检测通过并展示文件清单', '检测通过' in r.text and 'submit_张三.xlsx' in r.text and 'submit_李四.xlsx' in r.text)

r = s.post(BASE + f'/admin/theme/{tid3}/merge-excel', data={})
check('M2.2 执行合并返回 xlsx', r.status_code == 200 and r.content[:2] == b'PK', f'code={r.status_code}')
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
header = [ws.cell(row=1, column=c).value for c in range(1, 6)]
check('M2.3 表头来自首个文件', header == ['项目', '数量', '单价', '金额', '备注'], f'header={header}')
rows = [[ws.cell(row=r, column=c).value for c in range(1, 6)] for r in range(2, ws.max_row + 1)]
check('M2.4 按行堆叠 5 条数据', len(rows) == 5, f'rows={rows}')
check('M2.5 数据内容完整（含跨文件顺序）', rows[0][0] == '设备A' and rows[3][0] == '设备A' and rows[4][0] == '设备D', f'rows={rows}')
imgs = ws._images
check('M2.6 两张图片都保留', len(imgs) == 2, f'n={len(imgs)}')
anchors_ok = all(isinstance(im.anchor, TwoCellAnchor) for im in imgs)
check('M2.7 图片统一嵌入单元格（TwoCellAnchor）', anchors_ok)
# 图片应随各自数据行下沉：张三的图锚在第 2 行（索引1），李四的图锚在第 5 行（索引4）
img_rows = sorted(im.anchor._from.row for im in imgs)
check('M2.8 图片跟随数据行下沉（行 1 与 4）', img_rows == [1, 4], f'rows={img_rows}')
check('M2.9 表头样式保留（加粗）', ws.cell(row=1, column=1).font.bold is True)
check('M2.10 列宽沿用源表（14）', abs((ws.column_dimensions['A'].width or 0) - 14) < 0.01, f'w={ws.column_dimensions["A"].width}')
check('M2.11 工作表名沿用源表', ws.title == '提交表', f'title={ws.title}')

print()
print('=' * 70)
print('M3. 合并选项：重复行去重（相邻 + 跨文件全局）+ 标记来源')
print('=' * 70)

tid5, oids5 = create_theme('去重合并主题', ['张三', '李四'])
upload(oids5['张三'], 'submit_张三2.xlsx')
upload(oids5['李四'], 'submit_李四2.xlsx')

r = s.post(BASE + f'/admin/theme/{tid5}/merge-excel', data={})
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
n_plain = ws.max_row - 1
check('M3.0 不去重时 5 行', n_plain == 5, f'n={n_plain}')

r = s.post(BASE + f'/admin/theme/{tid5}/merge-excel', data={'merge_identical': 'on', 'mark_source': 'on'})
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
header = [ws.cell(row=1, column=c).value for c in range(1, ws.max_column + 1)]
check('M3.1 末尾新增“收集对象”列', header[-1] == '收集对象', f'header={header}')
rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)] for r in range(2, ws.max_row + 1)]
check('M3.2 相邻相同行被合并去重（5 → 4）', len(rows) == 4, f'rows={rows}')
srcs = [row[-1] for row in rows]
check('M3.3 来源列标注正确', srcs == ['张三', '张三', '张三 / 李四', '李四'], f'srcs={srcs}')
check('M3.4 被丢弃行的图片固定到保留行（仍 2 张图）', len(ws._images) == 2, f'n={len(ws._images)}')
img_rows = sorted(im.anchor._from.row for im in ws._images)
check('M3.5 去重后图片不错位（0-based 行 1 与 3）', img_rows == [1, 3], f'rows={img_rows}')

# 修复验证：非相邻的重复行现在也会被全局去重（原实现仅比较相邻行，
# 与“跨所有文件”的 UI 文案不符；tid3 的两个文件中“设备A”不相邻）
r = s.post(BASE + f'/admin/theme/{tid3}/merge-excel', data={'merge_identical': 'on', 'mark_source': 'on'})
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
n_cross = ws.max_row - 1
check('M3.6 修复验证：非相邻重复行全局去重（5 → 4）', n_cross == 4, f'n={n_cross}')
names_cross = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
check('M3.7 修复验证：保留首次出现的行（设备A 只出现一次）',
      names_cross == ['设备A', '设备B', '设备C', '设备D'], f'names={names_cross}')
srcs_cross = [ws.cell(row=r, column=ws.max_column).value for r in range(2, ws.max_row + 1)]
check('M3.8 修复验证：全局合并后来源标注正确',
      srcs_cross == ['张三 / 李四', '张三', '张三', '李四'], f'srcs={srcs_cross}')

print()
print('=' * 70)
print('M4. WPS 单元格内嵌截图（DISPIMG）保留')
print('=' * 70)

# 构造一个 WPS 风格的源文件：5 列结构 + DISPIMG 公式 + cellimages.xml
def make_wps_style_xlsx():
    wb = Workbook()
    ws = wb.active
    ws.title = '提交表'
    for c, h in enumerate(['项目', '数量', '单价', '金额', '备注'], 1):
        ws.cell(row=1, column=c, value=h)
    data = [['设备F', 1, 10.0, 10.0, 'wps'], ['设备G', 2, 20.0, 40.0, '']]
    for r, row in enumerate(data, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    # 第 2 行（首条数据行）放 DISPIMG 公式
    ws.cell(row=2, column=2, value='=_xlfn.DISPIMG("ID_1",1)')
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    pic = ('<etc:cellImage xmlns:xdr="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing" '
           'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
           'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
           '<xdr:pic><xdr:nvPicPr><xdr:cNvPr id="2" name="ID_1"/><xdr:cNvPicPr/></xdr:nvPicPr>'
           '<xdr:blipFill><a:blip r:embed="rId1"/><a:stretch><a:fillRect/></a:stretch></xdr:blipFill>'
           '<xdr:spPr><a:xfrm><a:off x="0" y="0"/><a:ext cx="952500" cy="666000"/></a:xfrm>'
           '<a:prstGeom prst="rect"><a:avLst/></a:prstGeom></xdr:spPr></xdr:pic></etc:cellImage>')
    images = [{'name': 'ID_1', 'media': fx('test_img.png'), 'ext': '.png', 'pic': pic}]
    return _inject_wps_cell_images(buf, images).read()

wps_bytes = make_wps_style_xlsx()
# 验证源文件确实是 WPS 风格（含 cellimages.xml）
zf = zipfile.ZipFile(io.BytesIO(wps_bytes))
check('M4.1 源文件含 cellimages.xml', 'xl/cellimages.xml' in zf.namelist())

tid4, oids4 = create_theme('WPS截图合并主题', ['张三', '王五'])
upload(oids4['张三'], 'submit_张三.xlsx')
upload(oids4['王五'], 'wps_截图.xlsx', content=wps_bytes)
r = s.get(BASE + f'/admin/theme/{tid4}/merge-excel')
check('M4.2 WPS 文件可参与合并检测', '检测通过' in r.text, r.text[:300])

r = s.post(BASE + f'/admin/theme/{tid4}/merge-excel', data={'mark_source': 'on'})
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)] for r in range(2, ws.max_row + 1)]
check('M4.3 数据行堆叠正确（3+2=5）', len(rows) == 5 and rows[3][0] == '设备F', f'rows={rows}')
check('M4.4 DISPIMG 公式已清除', all(not (isinstance(v, str) and v.startswith('=')) for row in rows for v in row), f'rows={rows}')
imgs = ws._images
check('M4.5 WPS 截图转换为真实图片（共 2 张）', len(imgs) == 2, f'n={len(imgs)}')
# 王五文件的截图在第 2 行（索引1）第 2 列（索引1），合并后应落到第 5 行（索引4）第 2 列
img_rows = sorted(im.anchor._from.row for im in imgs)
check('M4.6 截图跟随数据行下沉（0-based 行 1 与 4）', img_rows == [1, 4], f'rows={img_rows}')
img_cols = sorted(im.anchor._from.col for im in imgs)
check('M4.7 截图锚定列正确（WPS 图第 2 列、张三图第 4 列）', img_cols == [1, 3], f'cols={img_cols}')

print()
print('=' * 70)
print('M5. 合并各收集对象的数据（首列键互补合并：完整合并每个人填的数据）')
print('=' * 70)

def build_xlsx_bytes(rows, header=('项目', '数量', '单价', '金额', '备注')):
    wb = Workbook()
    ws = wb.active
    ws.title = '提交表'
    for c, h in enumerate(header, 1):
        ws.cell(row=1, column=c, value=h)
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            ws.cell(row=r, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()

tid6, oids6 = create_theme('互补合并主题', ['张三', '李四'])
upload(oids6['张三'], 'comp_张三.xlsx', content=build_xlsx_bytes([
    ['设备A', 2, None, None, '张三填的数量'],
    ['设备B', 1, 100.0, 100.0, ''],
]))
upload(oids6['李四'], 'comp_李四.xlsx', content=build_xlsx_bytes([
    ['设备A', None, 1500.5, 3001.0, '李四填的单价金额'],
    ['设备C', 5, 99.0, 495.0, ''],
]))

# 不勾选：两个 设备A 各成一行（共 4 行）
r = s.post(BASE + f'/admin/theme/{tid6}/merge-excel', data={})
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
n_raw = ws.max_row - 1
check('M5.0 不合并时 4 行（两个 设备A 各一行）', n_raw == 4, f'n={n_raw}')

# 勾选：设备A 互补合并为 1 行，设备B/设备C 各自保留 -> 3 行
r = s.post(BASE + f'/admin/theme/{tid6}/merge-excel',
           data={'merge_identical': 'on', 'mark_source': 'on'})
wb = load_workbook(io.BytesIO(r.content))
ws = wb.active
rows = [[ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1)]
check('M5.1 互补合并为 3 行（设备A 合并，设备B/设备C 保留）', len(rows) == 3, f'rows={rows}')
a = next(row for row in rows if row[0] == '设备A')
check('M5.2 设备A 数量补全自张三(2)', a[1] == 2, f'a={a}')
check('M5.3 设备A 单价补全自李四(1500.5)', a[2] == 1500.5, f'a={a}')
check('M5.4 设备A 金额补全自李四(3001.0)', a[3] == 3001.0, f'a={a}')
check('M5.5 冲突单元格保留先填写方(张三)的备注', a[4] == '张三填的数量', f'a={a}')
check('M5.6 设备A 来源列标注“张三 / 李四”', a[-1] == '张三 / 李四', f'a={a}')
check('M5.7 设备B/设备C 仍完整保留',
      any(row[0] == '设备B' and row[1] == 1 for row in rows)
      and any(row[0] == '设备C' and row[1] == 5 for row in rows), f'rows={rows}')

print()
print('=' * 70)
print(f'测试结果：通过 {len(PASS)} 项，失败 {len(FAIL)} 项')
if FAIL:
    print('失败项：')
    for name, detail in FAIL:
        print(f'  - {name}: {detail}')
print('=' * 70)
sys.exit(1 if FAIL else 0)
