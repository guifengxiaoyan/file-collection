# -*- coding: utf-8 -*-
"""生成测试用 fixture 文件：收集对象名单 Excel、带图表的提交 Excel、普通文件等。"""
import os
from io import BytesIO

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from PIL import Image as PILImage

FIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'fixtures')
os.makedirs(FIX, exist_ok=True)

# ---------- 1. 收集对象名单（Excel 导入用） ----------
wb = Workbook()
ws = wb.active
ws.title = '名单'
names = ['张三', '李四', '王五', '赵六', '钱七']
ws.cell(row=1, column=1, value='姓名')
for i, n in enumerate(names, 2):
    ws.cell(row=i, column=1, value=n)
wb.save(os.path.join(FIX, 'objects.xlsx'))

# ---------- 2. 生成一张测试图片 ----------
img = PILImage.new('RGB', (120, 80), (70, 130, 180))
px = img.load()
for x in range(120):
    for y in range(80):
        if (x // 20 + y // 20) % 2 == 0:
            px[x, y] = (240, 200, 80)
img_path = os.path.join(FIX, 'test_img.png')
img.save(img_path)

img2 = PILImage.new('RGB', (100, 100), (200, 80, 80))
img2_path = os.path.join(FIX, 'test_img2.png')
img2.save(img2_path)

# ---------- 3. 结构化提交 Excel（带样式、列宽、行高、浮动图片） ----------
def make_submit_excel(path, person, rows, img=None, anchor_col=3):
    wb = Workbook()
    ws = wb.active
    ws.title = '提交表'
    header = ['项目', '数量', '单价', '金额', '备注']
    header_font = Font(bold=True, color='FFFFFF', size=12)
    header_fill = PatternFill('solid', fgColor='4A6FA5')
    center = Alignment(horizontal='center', vertical='center')
    thin = Side(style='thin', color='999999')
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(header, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border
    ws.row_dimensions[1].height = 24
    for c in range(1, 6):
        ws.column_dimensions[chr(64 + c)].width = 14
    for r, row in enumerate(rows, 2):
        for c, v in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=v)
            cell.border = border
            if c == 4:
                cell.number_format = '#,##0.00'
        ws.row_dimensions[r].height = 30
    if img:
        xlimg = XLImage(img)
        xlimg.width, xlimg.height = 100, 66
        # 浮动图片（OneCellAnchor + ext），锚定在第 2 行
        xlimg.anchor = OneCellAnchor(
            _from=AnchorMarker(col=anchor_col, row=1, colOff=0, rowOff=0),
            ext=__import__('openpyxl.drawing.spreadsheet_drawing', fromlist=['XDRPositiveSize2D']).XDRPositiveSize2D(
                int(100 * 9525), int(66 * 9525)))
        ws.add_image(xlimg)
    wb.save(path)

rows_a = [
    ['设备A', 2, 1500.5, 3001.0, '含税'],
    ['设备B', 1, 899.0, 899.0, '不含税'],
    ['设备C', 3, 250.0, 750.0, ''],
]
rows_b = [
    ['设备A', 2, 1500.5, 3001.0, '含税'],   # 与 A 完全相同 -> 用于去重测试
    ['设备D', 5, 99.9, 499.5, '加急'],
]
make_submit_excel(os.path.join(FIX, 'submit_张三.xlsx'), '张三', rows_a, img=img_path)
make_submit_excel(os.path.join(FIX, 'submit_李四.xlsx'), '李四', rows_b, img=img2_path, anchor_col=2)

# ---------- 3b. 相邻重复行 fixture（验证“合并相同行”去重） ----------
rows_a2 = [
    ['设备A', 2, 1500.5, 3001.0, '含税'],
    ['设备B', 1, 899.0, 899.0, '不含税'],
    ['设备C', 3, 250.0, 750.0, ''],
]
rows_b2 = [
    ['设备C', 3, 250.0, 750.0, ''],   # 与张三最后一行完全相同且相邻 -> 应被去重
    ['设备D', 5, 99.9, 499.5, '加急'],
]
make_submit_excel(os.path.join(FIX, 'submit_张三2.xlsx'), '张三', rows_a2, img=img_path)
make_submit_excel(os.path.join(FIX, 'submit_李四2.xlsx'), '李四', rows_b2, img=img2_path, anchor_col=2)

# 列数不一致的文件（6 列）——用于“格式不一致”检测
make_submit_excel(os.path.join(FIX, 'submit_王五_6cols.xlsx'), '王五',
                  [['设备E', 1, 2, 3, 'x', 'y']] + [[None] * 6])

# ---------- 4. 普通文件 ----------
with open(os.path.join(FIX, '说明.txt'), 'w', encoding='utf-8') as f:
    f.write('测试附件内容\n')

import zipfile
with zipfile.ZipFile(os.path.join(FIX, '打包.zip'), 'w') as zf:
    zf.writestr('readme.txt', 'hello')

# 模拟 PDF（内容无所谓，后缀合法即可）
with open(os.path.join(FIX, '报告.pdf'), 'wb') as f:
    f.write(b'%PDF-1.4 fake pdf for test')

# 不合法扩展名
with open(os.path.join(FIX, '病毒.exe'), 'wb') as f:
    f.write(b'MZ fake exe')

# ---------- 5. 带中文/特殊字符文件名的文件 ----------
with open(os.path.join(FIX, '汇总表 2026（最终版）.xlsx'), 'wb') as f:
    f.write(b'fake xlsx bytes')

print('fixtures created:', sorted(os.listdir(FIX)))
